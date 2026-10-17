import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from .base_parser import BaseParser


class DocxParser(BaseParser):
    """Extract paragraphs and Word tables in document order.

    Structured resume sections are frequently stored in tables. Parsing only
    Document.paragraphs silently drops their contents while keeping the heading.
    """

    @staticmethod
    def _table_lines(table: Table) -> list[str]:
        lines: list[str] = []
        for row in table.rows:
            cells: list[str] = []
            for cell in row.cells:
                value = " ".join(
                    paragraph.text.strip()
                    for paragraph in cell.paragraphs
                    if paragraph.text.strip()
                ).strip()
                # Merged cells can repeat in python-docx's row cell collection.
                if value and (not cells or value != cells[-1]):
                    cells.append(value)
            if cells:
                lines.append(" | ".join(cells))
        return lines

    def parse(self, file_path: str) -> str:
        try:
            document = docx.Document(file_path)
            lines: list[str] = []
            # The body XML preserves paragraph/table order as seen by readers.
            for child in document.element.body.iterchildren():
                if child.tag == qn("w:p"):
                    value = Paragraph(child, document).text.strip()
                    if value:
                        lines.append(value)
                elif child.tag == qn("w:tbl"):
                    lines.extend(self._table_lines(Table(child, document)))
            return "\n".join(lines)
        except Exception as exc:
            raise RuntimeError(f"Error parsing DOCX {file_path}: {exc}") from exc
