from docx import Document as WordDocument

from app.parsers.docx_parser import DocxParser


def test_docx_parser_extracts_table_rows_in_document_order(tmp_path):
    source = WordDocument()
    source.add_paragraph("PROJECTS UNDERTAKEN")
    source.add_paragraph("An example project")
    source.add_paragraph("EDUCATION DETAILS")
    table = source.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "Course"
    table.rows[0].cells[1].text = "Institute"
    table.rows[0].cells[2].text = "Score"
    row = table.add_row().cells
    row[0].text = "UG"
    row[1].text = "Vishwakarma Institute of Information Technology (VIIT), Pune"
    row[2].text = "8.3 / 10 CGPA"
    source.add_paragraph("CERTIFICATIONS")
    source.add_paragraph("Full Stack Software Developer Certificate")

    path = tmp_path / "resume.docx"
    source.save(path)
    parsed = DocxParser().parse(str(path))

    assert parsed.index("PROJECTS UNDERTAKEN") < parsed.index("EDUCATION DETAILS")
    assert parsed.index("EDUCATION DETAILS") < parsed.index("Course | Institute | Score")
    assert parsed.index("Course | Institute | Score") < parsed.index("CERTIFICATIONS")
    assert "Vishwakarma Institute of Information Technology (VIIT), Pune" in parsed
    assert "8.3 / 10 CGPA" in parsed
